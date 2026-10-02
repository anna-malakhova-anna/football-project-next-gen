import collections
import hashlib

import pyarrow.parquet as pq
import pytest
from openpyxl import load_workbook

from generator import main
from generator.population import load_config
from generator.sources import intl_federation as intl


@pytest.fixture(scope="module")
def run(tmp_path_factory):
    base = tmp_path_factory.mktemp("gen_intl")
    cfg = load_config()
    res = main.run(cfg, raw_dir=str(base / "raw"), truth_dir=str(base / "truth"))
    out = next(o for o in res["outputs"] if o.source_system == "INTL_FED")
    return cfg, base, out


def _sheets(base):
    """-> {(year, sheet_name): list of raw rows (tuples)}"""
    res = {}
    for f in sorted((base / "raw" / "intl").glob("federation_returns_*.xlsx")):
        year = int(f.stem.rsplit("_", 1)[1])
        wb = load_workbook(f)
        for ws in wb:
            res[(year, ws.title)] = list(ws.iter_rows(values_only=True))
    return res


def _data_rows(rows):
    header = tuple(intl.COLUMNS)
    start = rows.index(header) + 1
    body = rows[start:]
    return [r for r in body if r[1] is not None]  # drops the "Total: N" row


def test_one_workbook_per_season_one_sheet_per_country(run):
    cfg, base, _ = run
    files = sorted(p.name for p in (base / "raw" / "intl").glob("*.xlsx"))
    assert files == [f"federation_returns_{y}.xlsx" for y in cfg["seasons"]["years"]]
    names = {name for (_, name) in _sheets(base)}
    assert names == set(intl.SHEET_NAME.values())


def test_title_rows_and_summary_rows(run):
    cfg, base, _ = run
    sc = cfg["sources"]["intl"]
    for (year, name), rows in _sheets(base).items():
        code = next(c for c, n in intl.SHEET_NAME.items() if n == name)
        if code in sc["sheets_with_title_rows"]:
            assert rows[0][0].startswith("Federation Return") and all(v is None for v in rows[1])
            assert rows[2] == tuple(intl.COLUMNS)
        else:
            assert rows[0] == tuple(intl.COLUMNS)
        last = rows[-1]
        if code in sc["sheets_with_summary_row"]:
            assert last[0] == f"Total: {len(_data_rows(rows))}" and all(v is None for v in last[1:])
        else:
            assert not str(last[0]).startswith("Total")


def test_uk_dob_excel_serials(run):
    cfg, base, _ = run
    uk = [r for (y, n), rows in _sheets(base).items() if n == "United Kingdom" for r in _data_rows(rows)]
    ser = sum(isinstance(r[2], (int, float)) for r in uk)
    assert ser / len(uk) == pytest.approx(cfg["defects"]["intl_federation"]["uk_dob_excel_serial"], abs=0.01)
    others = [r for (y, n), rows in _sheets(base).items() if n != "United Kingdom" for r in _data_rows(rows)]
    assert all(isinstance(r[2], str) and len(r[2]) == 10 and r[2][2] == "/" for r in others)


def test_country_inconsistency_and_label_variety(run):
    cfg, base, _ = run
    rows = [r for rs in _sheets(base).values() for r in _data_rows(rs)]
    base_forms = set(intl.ISO3.values())
    off = sum(r[6] not in base_forms for r in rows) / len(rows)
    assert off == pytest.approx(cfg["defects"]["intl_federation"]["country_value_inconsistency"], abs=0.02)
    assert {"MEX", "Mexico", "MX"} <= {r[6] for r in rows}
    assert {r[3] for r in rows} >= {"W", "M", "F", "Mujer", "Hombre"}
    assert {r[8] for r in rows} == {"Yes", "No", None}


def test_japan_family_name_first_and_panama_ids_reissued(run):
    _, base, out = run
    reg = {p["true_person_id"]: p for p in pq.read_table(base / "truth" / "person_registry.parquet").to_pylist()}
    truth = {t["source_record_id"]: t["true_person_id"] for t in out.truth}
    ids_by_person = collections.defaultdict(set)
    for (year, name), rows in _sheets(base).items():
        for r in _data_rows(rows):
            p = reg[truth[f"{r[0]}|{year}"]]
            ids_by_person[(name, p["true_person_id"])].add(r[0])
            if name == "Japan":
                assert r[1] == f"{p['family_name']} {p['given_name']}"
            else:
                assert r[1] == f"{p['given_name']} {p['family_name']}"
    multi = [(k, v) for k, v in ids_by_person.items()]
    # Panama: a returning player gets a new id each year; all other countries keep one
    for (name, _), ids in multi:
        if name != "Panama":
            assert len(ids) == 1
    assert any(len(ids) > 1 for (name, _), ids in multi if name == "Panama")


def test_player_ids_unique_within_sheet_and_truth_complete(run):
    _, base, out = run
    keys = set()
    for (year, name), rows in _sheets(base).items():
        data = _data_rows(rows)
        ids = [r[0] for r in data]
        assert len(ids) == len(set(ids))
        keys |= {f"{i}|{year}" for i in ids}
    assert keys == {t["source_record_id"] for t in out.truth}
    assert out.n_records == len(keys)


def test_categories_known_and_minor_email_mostly_blank(run):
    _, base, _ = run
    rows = [r for rs in _sheets(base).values() for r in _data_rows(rs)]
    allowed = {l for k in intl.CATEGORY_LABELS.values() for l in k.values()}
    assert {r[4] for r in rows} <= allowed
    youth = [r for r in rows if r[4] in {"U8", "U10", "U12", "U15", "Juniors", "Juvenil", "Junioren"}]
    assert sum(r[7] is None for r in youth) / len(youth) > 0.7


def test_junk_rows_in_manifest(run):
    _, _, out = run
    counts = collections.Counter(m["defect_id"] for m in out.manifest)
    assert counts["junk_title_row"] == counts["junk_blank_row"] == 2 * 3   # DE+AT x 3 years
    assert counts["junk_summary_row"] == 4 * 3                              # MX, GB, DE, PA x 3 years


def test_xlsx_deterministic(run, tmp_path):
    cfg, _, _ = run
    hashes = []
    for i in range(2):
        main.run(cfg, raw_dir=str(tmp_path / f"r{i}"), truth_dir=str(tmp_path / f"t{i}"))
        hashes.append({p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (tmp_path / f"r{i}" / "intl").glob("*.xlsx")})
    assert hashes[0] == hashes[1] and len(hashes[0]) == 3
