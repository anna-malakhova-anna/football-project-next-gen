"""Land the generated source files in DuckDB schema `raw` with NO transformations.

All values are stored as the strings found in the files (empty CSV fields stay '', Excel serials
stay their raw numeric string). Typing, trimming, NULLIF and junk removal happen in dbt staging.
Ground truth (data/ground_truth/) is never read here.

  raw.org_registry              one row per known organization and season: where it is, its channel, and whether it sent a file (the coverage denominator)
  raw.us_league_registrations   one row per CSV data row; `_row_number` = 1-based data row index
  raw.ca_club_records           one row per API record; nested JSON kept whole in `record_json`
  raw.school_headcounts         one row per CSV data row of the aggregate school / recreation program headcounts (no individuals)
  raw.fan_accounts              one row per account of the fan app / ticketing snapshot (CSV; adults only)
  raw.fan_consent_events        one row per consent change line (JSON Lines), kept whole in `record_json`
  raw.fan_transactions          one row per transaction line (JSON Lines, one file per month), kept whole in `record_json`
  raw.fan_events                one row per delivered event line (JSON Lines, one file per delivery month); the line is kept whole in `record_json`
  raw.intl_federation_returns   every worksheet row, header/title/blank/summary rows included, as
                                positional columns c1..cN; `_excel_row` = 1-based worksheet row.
                                Staging finds the header row and drops the junk.

Idempotent: tables are CREATE OR REPLACE'd on every run.
"""
from __future__ import annotations

import csv
import datetime as dt
import json
from pathlib import Path

import duckdb
import pyarrow as pa
from openpyxl import load_workbook

from generator.population import load_config

RAW_TABLES = ("org_registry", "us_league_registrations", "ca_club_records", "intl_federation_returns", "school_headcounts", "fan_accounts", "fan_events",
              "fan_consent_events", "fan_transactions")


def _rel(path: Path, raw: Path) -> str:
    return path.relative_to(raw).as_posix()


def read_us(raw: Path) -> tuple[list[str], list[list]]:
    files = sorted((raw / "us_league").glob("*.csv"))
    columns = None
    rows: list[list] = []
    for f in files:
        with open(f, newline="", encoding="utf-8") as fh:
            reader = csv.reader(fh)
            header = next(reader)
            columns = columns or header
            if header != columns:
                raise ValueError(f"{f.name}: header differs from the other US files")
            for n, rec in enumerate(reader, start=1):
                rows.append(rec + [_rel(f, raw), n])
    return columns, rows


def read_org_registry(raw: Path) -> tuple[list[str], list[list]]:
    """The organization registry (what the program knows): one row per known org and season, with whether it sent a file."""
    path = raw / "org_registry.csv"
    with open(path, newline="", encoding="utf-8") as fh:
        reader = csv.reader(fh)
        header = next(reader)
        return header, [rec + [_rel(path, raw), n] for n, rec in enumerate(reader, start=1)]


def read_ca(raw: Path) -> list[list]:
    rows = []
    for f in sorted((raw / "ca_club").glob("page_*.json")):
        page = json.loads(f.read_text(encoding="utf-8"))
        for rec in page["records"]:
            rows.append([json.dumps(rec, ensure_ascii=False), page["page"], _rel(f, raw)])
    return rows


def read_school(raw: Path) -> tuple[list[str], list[list]]:
    """Aggregate headcounts: one row per CSV data row, no individuals."""
    files = sorted((raw / "school").glob("*.csv"))
    columns = None
    rows: list[list] = []
    for f in files:
        with open(f, newline="", encoding="utf-8") as fh:
            reader = csv.reader(fh)
            header = next(reader)
            columns = columns or header
            if header != columns:
                raise ValueError(f"{f.name}: header differs from the other school files")
            for n, rec in enumerate(reader, start=1):
                rows.append(rec + [_rel(f, raw), n])
    return columns, rows


def read_fan_accounts(raw: Path) -> tuple[list[str], list[list]]:
    path = raw / "fan" / "fan_accounts.csv"
    with open(path, newline="", encoding="utf-8") as fh:
        reader = csv.reader(fh)
        header = next(reader)
        return header, [rec + [_rel(path, raw), n] for n, rec in enumerate(reader, start=1)]


def _read_jsonl(raw: Path, pattern: str) -> list[list]:
    rows = []
    for f in sorted((raw / "fan").glob(pattern)):
        with open(f, encoding="utf-8") as fh:
            for n, line in enumerate(fh, start=1):
                rows.append([line.rstrip("\n"), _rel(f, raw), n])
    return rows


def read_fan_events(raw: Path) -> list[list]:
    """One row per event line, kept whole. Files are delivery months, so a late event sits in a later file than its event_ts."""
    rows = []
    for f in sorted((raw / "fan").glob("events_*.jsonl")):
        with open(f, encoding="utf-8") as fh:
            for n, line in enumerate(fh, start=1):
                rows.append([line.rstrip("\n"), _rel(f, raw), n])
    return rows


def read_intl(raw: Path) -> tuple[int, list[list]]:
    rows = []
    width = 0
    for f in sorted((raw / "intl").glob("*.xlsx")):
        wb = load_workbook(f, data_only=True)
        for ws in wb:
            for n, cells in enumerate(ws.iter_rows(values_only=True), start=1):
                vals = [None if v is None else str(v) for v in cells]
                width = max(width, len(vals))
                rows.append([vals, ws.title, _rel(f, raw), n])
    return width, [vals + [None] * (width - len(vals)) + [sheet, src, n] for vals, sheet, src, n in rows]


def _table(columns: list[str], rows: list[list], types: dict[str, pa.DataType]) -> pa.Table:
    arrays = [pa.array([r[i] for r in rows], type=types.get(c, pa.string())) for i, c in enumerate(columns)]
    return pa.Table.from_arrays(arrays, names=columns)


def load(raw_dir: str | None = None, db_path: str | None = None, cfg: dict | None = None) -> dict[str, int]:
    cfg = cfg or load_config()
    raw = Path(raw_dir or cfg["paths"]["raw_dir"])
    db = Path(db_path or cfg["paths"]["warehouse"])
    db.parent.mkdir(parents=True, exist_ok=True)
    loaded_at = dt.datetime.now().replace(microsecond=0)

    reg_cols, reg_rows = read_org_registry(raw)
    us_cols, us_rows = read_us(raw)
    ca_rows = read_ca(raw)
    width, intl_rows = read_intl(raw)
    school_cols, school_rows = read_school(raw)
    fan_cols, fan_acc_rows = read_fan_accounts(raw)
    fan_event_rows = read_fan_events(raw)
    fan_consent_rows = _read_jsonl(raw, "fan_consent_events.jsonl")
    fan_txn_rows = _read_jsonl(raw, "fan_transactions_*.jsonl")

    tables = {
        "org_registry": _table(reg_cols + ["_source_file", "_row_number"], reg_rows, {"_row_number": pa.int32()}),
        "us_league_registrations": _table(
            us_cols + ["_source_file", "_row_number"], us_rows, {"_row_number": pa.int32()}),
        "ca_club_records": _table(
            ["record_json", "_page", "_source_file"], ca_rows, {"_page": pa.int32()}),
        "intl_federation_returns": _table(
            [f"c{i}" for i in range(1, width + 1)] + ["_sheet_name", "_source_file", "_excel_row"],
            intl_rows, {"_excel_row": pa.int32()}),
        "school_headcounts": _table(school_cols + ["_source_file", "_row_number"], school_rows, {"_row_number": pa.int32()}),
        "fan_accounts": _table(fan_cols + ["_source_file", "_row_number"], fan_acc_rows, {"_row_number": pa.int32()}),
        "fan_events": _table(["record_json", "_source_file", "_line_number"], fan_event_rows, {"_line_number": pa.int32()}),
        "fan_consent_events": _table(["record_json", "_source_file", "_line_number"], fan_consent_rows, {"_line_number": pa.int32()}),
        "fan_transactions": _table(["record_json", "_source_file", "_line_number"], fan_txn_rows, {"_line_number": pa.int32()}),
    }
    con = duckdb.connect(str(db))
    try:
        con.execute("CREATE SCHEMA IF NOT EXISTS raw")
        for name, tbl in tables.items():
            con.register("_landing", tbl)
            con.execute(f"CREATE OR REPLACE TABLE raw.{name} AS SELECT * FROM _landing")
            con.unregister("_landing")
            con.execute(f"ALTER TABLE raw.{name} ADD COLUMN _loaded_at TIMESTAMP DEFAULT TIMESTAMP '{loaded_at}'")
            con.execute(f"UPDATE raw.{name} SET _loaded_at = TIMESTAMP '{loaded_at}'")
        con.execute("ALTER TABLE raw.ca_club_records ALTER record_json SET DATA TYPE JSON")
        for t in ("fan_events", "fan_consent_events", "fan_transactions"):
            con.execute(f"ALTER TABLE raw.{t} ALTER record_json SET DATA TYPE JSON")
        counts = {n: con.execute(f"SELECT count(*) FROM raw.{n}").fetchone()[0] for n in tables}
    finally:
        con.close()
    expected = {"org_registry": len(reg_rows), "us_league_registrations": len(us_rows), "ca_club_records": len(ca_rows),
                "intl_federation_returns": len(intl_rows), "school_headcounts": len(school_rows),
                "fan_accounts": len(fan_acc_rows), "fan_events": len(fan_event_rows),
                "fan_consent_events": len(fan_consent_rows), "fan_transactions": len(fan_txn_rows)}
    if counts != expected:
        raise RuntimeError(f"row count mismatch: loaded {counts}, files have {expected}")
    return counts


if __name__ == "__main__":
    for table, n in load().items():
        print(f"raw.{table}: {n} rows")
